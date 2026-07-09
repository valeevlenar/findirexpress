from datetime import datetime
import logging
import io
from openpyxl.drawing.image import Image as ExcelImage
import matplotlib.pyplot as plt
import matplotlib.patches as patches
from openpyxl.styles import Font, PatternFill, Border, Side
from app.admin.admin_message import send_message_to_admin
from app.database.models import async_session
from app.database.support_functions import get_seller_inn_by_seller_id, get_company_name_by_seller_id
from app.wrappers import log_and_notify_admin, with_session


# --- Расчет разниц ---
class Change:
    @staticmethod
    def change(value: any, previous: any):
        return value - previous

    @staticmethod
    def change_percent(value: any, previous: any):
        try:
            if previous != 0:
                change_percent = (value / previous - 1)
            else:
                change_percent = 0.0 if value == 0 else 1.0  # Обработка нулевого предыдущего значения

        except:
            change_percent = 0.0

        return change_percent

# --- Formatters ---
class FormatterWithUnits:
    @classmethod
    def __getitem__(cls, key):
        if key == "currency":
            return cls.currency
        elif key == "percent":
            return cls.percent
        elif key == "units":
            return cls.units
        elif key == "days":
            return cls.days
        else:
            raise KeyError(f"Unknown formatter type: {key}")

    @staticmethod
    def currency(value: float) -> str:
        return f"{value:,.0f}".replace(",", " ") + " р."

    @staticmethod
    def percent(value: float) -> str:
        return f"{value:.1%}"

    @staticmethod
    def units(value: float) -> str:
        return f"{value:,.0f}".replace(",", " ") + " шт."

    @staticmethod
    def days(value: float) -> str:
        return f"{value:,.0f}".replace(",", " ") + " дн."

class Formatter:
    @classmethod
    def __getitem__(cls, key):
        if key == "currency":
            return cls.currency
        elif key == "percent":
            return cls.percent
        elif key == "units":
            return cls.units
        elif key == "days":
            return cls.days
        else:
            raise KeyError(f"Unknown formatter type: {key}")

    @staticmethod
    def currency(value: float) -> str:
        return f"{value:,.0f}".replace(",", " ")

    @staticmethod
    def percent(value: float) -> str:
        return f"{value:.1%}"

    @staticmethod
    def units(value: float) -> str:
        return f"{value:,.0f}".replace(",", " ")

    @staticmethod
    def days(value: float) -> str:
        return f"{value:,.0f}".replace(",", " ")

class Arrow:
    @staticmethod
    def get_arrow_and_color(arrow_type, change_value):
        arrow = '\u2191'   #'↑'
        change_color = 'green'
        if arrow_type == 'straight':
            if change_value >= 0.0:
                arrow = '\u2191'   #'↑'
                change_color = 'green'
            else:
                arrow = '\u2193' #'↓'
                change_color = 'red'
        elif arrow_type == 'back':
            if change_value >= 0.0:
                arrow = '\u2191'   #'↑'
                change_color = 'red'
            else:
                arrow = '\u2193' #'↓'
                change_color = 'green'
        return arrow, change_color

#заполняем эксель файл

@log_and_notify_admin
async def generate_dashboard_item(header, arrow_type, value_main, prev_main, main_value_type, value_second, prev_second, second_value_type):

    main_change = Change.change(value_main, prev_main)
    main_change_percent = Change.change_percent(value_main, prev_main)
    main_line = f'{getattr(FormatterWithUnits, main_value_type)(value_main)}'

    arrow, change_color = Arrow.get_arrow_and_color(arrow_type, main_change)
    change_line = f'{arrow} {getattr(Formatter, main_value_type)(main_change)} ({Formatter.percent(main_change_percent)})'

    second_change_line = None
    second_change_color = 'green'
    if value_second:
        second_change = Change.change(value_second, prev_second)
        second_change_percent = Change.change_percent(value_second, prev_second)
        second_arrow, second_change_color = Arrow.get_arrow_and_color(arrow_type, second_change)
        main_line += f' / {getattr(FormatterWithUnits,second_value_type)(value_second)}'
        second_change_line = f' / {second_arrow} {getattr(Formatter, second_value_type)(second_change)} ({Formatter.percent(second_change_percent)})'
    image_buffer = await create_image_for_dashboard(header, main_line, change_line, change_color,
                                                    second_change_line, second_change_color)
    img = ExcelImage(image_buffer)
    return img

async def create_image_for_dashboard (header, main_line, change_line, change_color, second_change_line, second_change_color):
    try:
        # Создаем буфер в памяти

        plt.rcParams['font.family'] = 'Tahoma'
        #plt.rcParams['font.sans-serif'] = ['Tahoma']
        plt.rcParams['font.family'] = ['Tahoma', 'DejaVu Sans', 'Arial Unicode MS', 'sans-serif'] # Добавляем запасные варианты
        plt.rcParams['axes.unicode_minus'] = False
        plt.rcParams['mathtext.default'] = 'regular'


        buf = io.BytesIO()
        fig, ax = plt.subplots(figsize=(4.95, 1.65))

        # Добавляем тень
        shadow = patches.FancyBboxPatch((0.03, 0.024), 0.94, 0.94, boxstyle="round,pad=0.03",
                                         linewidth=1, facecolor='#E7E7F1', edgecolor='#D5D5E7', alpha = 0.5)
        ax.add_patch(shadow)
        # Создание прямоугольника с закругленными углами
        rounded_rect = patches.FancyBboxPatch((0.024, 0.035), 0.94, 0.94, boxstyle="round,pad=0.03",
                                               linewidth=1, facecolor='white', edgecolor='#D5D5E7')

        ax.add_patch(rounded_rect)

        # Добавление текста

        ax.text(0.05, 0.75, header, fontsize=15, ha='left', fontdict=None)
        ax.text(0.05, 0.45, main_line, fontsize=17, ha='left',weight="bold")
        second_text = ax.text(0.05, 0.20, change_line, fontsize=14, ha='left', color = change_color)
        if second_change_line and second_change_color:
            second_text = ax.annotate(second_change_line, xycoords=second_text, xy=(1, 0.2), fontsize=14, ha='left', color=second_change_color)

        # Убираем оси
        ax.axis('off')

        # Сохранение изображения в буфер
        plt.savefig(buf, format='png', bbox_inches='tight', transparent=True)
        # Сохранение изображения
        plt.savefig('order_count.png', bbox_inches='tight', transparent=True)
        plt.close(fig)

        # Перемещаем указатель буфера в начало
        buf.seek(0)
        return buf
    except Exception as e:
        # Запись ошибки в лог
        await send_message_to_admin(f'Ошибка в функции создания картинки для дэшборда:'
                                    f'\nОшибка: {e}')
        logging.exception("An error occurred: %s", exc_info=e)

@with_session
async def create_dashboard(session, ws, seller_id, metrics, prev_metrics, date_from, date_to, prev_date_from, prev_date_to):
    try:
        seller_inn = await get_seller_inn_by_seller_id(session, seller_id)
        company_name = await get_company_name_by_seller_id(session, seller_id)
        company_plus_inn = f'{company_name}, ИНН {seller_inn}'
        ws = ws

        # Заполняем шапку
        ws['B2'] = company_plus_inn
        ws['E5'] = datetime.strftime(date_from, '%Y-%m-%d')
        ws['E6'] = datetime.strftime(date_to, '%Y-%m-%d')

        # Распределение показателей по блокам
        orders_anchor = 'B8'
        aop_anchor = 'H8'
        buyout_percent_anchor = 'N8'

        revenue_before_spp_anchor = 'B15'
        commission_before_spp_anchor = 'H15'
        spp_amount_anchor = 'N15'

        sales_after_spp_anchor = 'B22'
        commission_after_spp_anchor = 'H22'
        logistic_anchor = 'N22'

        profit_anchor = 'B29'
        profit_margin_anchor = 'H29'
        roi_anchor = 'N29'

        cost_of_sales_anchor = 'B36'
        marketing_anchor = 'H36'
        drr_anchor = 'N36'

        storage_anchor = 'B43'
        tax_costs_anchor = 'H43'
        other_deductions_anchor = 'N43'

        goods_at_cost_anchor = 'B50'
        goods_at_sell_price_anchor = 'H50'
        turnover_anchor = 'N50'

        # markup_anchor = 'H29'
        # slow_moving_anchor = 'B43'

        dashboard_items = []
        # Заказы
        orders_dashboard_item = {'header':"Заказы",
                                 'anchor': orders_anchor,
                                 'value_main':metrics.orders_sum,
                                 'prev_main': prev_metrics.orders_sum,
                                 'main_value_type':'currency',
                                 'value_second': metrics.orders_count,
                                 'prev_second': prev_metrics.orders_count,
                                 'second_value_type': 'units',
                                 'arrow_type': 'straight'}
        dashboard_items.append(orders_dashboard_item)

        # Продажи
        sales_dashboard_item = {'header': "Выручка (выкупы) после СПП",
                                 'anchor': sales_after_spp_anchor,
                                 'value_main': metrics.revenue,
                                 'prev_main': prev_metrics.revenue,
                                 'main_value_type': 'currency',
                                 'value_second': metrics.sales_count,
                                 'prev_second': prev_metrics.sales_count,
                                 'second_value_type': 'units',
                                 'arrow_type': 'straight'}
        dashboard_items.append(sales_dashboard_item)

        # Процент выкупа
        buyout_percent_dashboard_item = {'header': 'Процент выкупа',
                                         'anchor': buyout_percent_anchor,
                                         'value_main': metrics.buyout_percent,
                                         'prev_main': prev_metrics.buyout_percent,
                                         'main_value_type': 'percent',
                                         'value_second': None,
                                         'prev_second': None,
                                         'second_value_type': 'units',
                                         'arrow_type': 'straight'}
        dashboard_items.append(buyout_percent_dashboard_item)

        # Выручка до СПП
        revenue_before_spp_item = {'header': 'Выручка (выкупы) до СПП',
                                         'anchor': revenue_before_spp_anchor,
                                         'value_main': metrics.revenue_before_spp,
                                         'prev_main': prev_metrics.revenue_before_spp,
                                         'main_value_type': 'currency',
                                         'value_second': metrics.sales_count,
                                         'prev_second': prev_metrics.sales_count,
                                         'second_value_type': 'units',
                                         'arrow_type': 'straight'}
        dashboard_items.append(revenue_before_spp_item)

        # Комиссия до СПП
        commission_before_spp_item = {'header': 'Комиссия до СПП / % от выручки',
                                         'anchor': commission_before_spp_anchor,
                                         'value_main': metrics.commission_before_spp,
                                         'prev_main': prev_metrics.commission_before_spp,
                                         'main_value_type': 'currency',
                                         'value_second': metrics.commission_before_spp_rate,
                                         'prev_second': prev_metrics.commission_before_spp_rate,
                                         'second_value_type': 'percent',
                                         'arrow_type': 'back'}
        dashboard_items.append(commission_before_spp_item)

        # СПП
        spp_amount_item = {'header': 'СПП / % от выручки',
                                         'anchor': spp_amount_anchor,
                                         'value_main': metrics.spp_amount,
                                         'prev_main': prev_metrics.spp_amount,
                                         'main_value_type': 'currency',
                                         'value_second': metrics.spp_rate,
                                         'prev_second': prev_metrics.spp_rate,
                                         'second_value_type': 'percent',
                                         'arrow_type': 'straight'}
        dashboard_items.append(spp_amount_item)

        # Прибыль

        net_profit_dashboard_item = {'header': "Чистая прибыль",
                                         'anchor': profit_anchor,
                                         'value_main': metrics.net_profit,
                                         'prev_main': prev_metrics.net_profit,
                                         'main_value_type': 'currency',
                                         'value_second': None,
                                         'prev_second': None,
                                         'second_value_type': 'units',
                                         'arrow_type': 'straight'}
        dashboard_items.append(net_profit_dashboard_item)

        # Маржинальность

        margin_dashboard_item = {'header': "Маржинальность до / после СПП",
                                         'anchor': profit_margin_anchor,
                                         'value_main': metrics.margin_rate_before_spp,
                                         'prev_main': prev_metrics.margin_rate_before_spp,
                                         'main_value_type': 'percent',
                                         'value_second': metrics.margin_rate,
                                         'prev_second': prev_metrics.margin_rate,
                                         'second_value_type': 'percent',
                                         'arrow_type': 'straight'}
        dashboard_items.append(margin_dashboard_item)

        # "ROI"

        margin_dashboard_item = {'header': "ROI",
                                         'anchor': roi_anchor,
                                         'value_main': metrics.roi,
                                         'prev_main': prev_metrics.roi,
                                         'main_value_type': 'percent',
                                         'value_second': None,
                                         'prev_second': None,
                                         'second_value_type': 'units',
                                         'arrow_type': 'straight'}
        dashboard_items.append(margin_dashboard_item)

        # Себестоимость

        cost_of_sales_dashboard_item = {'header': "Себестоимость",
                                 'anchor': cost_of_sales_anchor,
                                 'value_main': -metrics.cost_of_sales,
                                 'prev_main': -prev_metrics.cost_of_sales,
                                 'main_value_type': 'currency',
                                 'value_second': None,
                                 'prev_second': None,
                                 'second_value_type': 'units',
                                 'arrow_type': 'back'}
        dashboard_items.append(cost_of_sales_dashboard_item)


        # Реклама

        dashboard_item = {'header': "Маркетинг",
                                        'anchor': marketing_anchor,
                                        'value_main': metrics.marketing,
                                        'prev_main': prev_metrics.marketing,
                                        'main_value_type': 'currency',
                                        'value_second': None,
                                        'prev_second': None,
                                        'second_value_type': 'percent',
                                        'arrow_type': 'back'}
        dashboard_items.append(dashboard_item)

        # Реклама

        dashboard_item = {'header': "ДРР до / после СПП",
                                        'anchor': drr_anchor,
                                        'value_main': metrics.drr_before_spp,
                                        'prev_main': prev_metrics.drr_before_spp,
                                        'main_value_type': 'percent',
                                        'value_second': metrics.drr,
                                        'prev_second': prev_metrics.drr,
                                        'second_value_type': 'percent',
                                        'arrow_type': 'back'}
        dashboard_items.append(dashboard_item)

        # Комиссия

        dashboard_item = {'header': "Комиссия после СПП / % от выручки",
                            'anchor': commission_after_spp_anchor,
                            'value_main': metrics.commission,
                            'prev_main': prev_metrics.commission,
                            'main_value_type': 'currency',
                            'value_second': metrics.commission_rate,
                            'prev_second': prev_metrics.commission_rate,
                            'second_value_type': 'percent',
                            'arrow_type': 'back'}
        dashboard_items.append(dashboard_item)

        # Средний чек

        dashboard_item = {'header': "Средний чек / Наценка",
                            'anchor': aop_anchor,
                            'value_main': metrics.aop,
                            'prev_main': prev_metrics.aop,
                            'main_value_type': 'currency',
                            'value_second': metrics.markup,
                            'prev_second': prev_metrics.markup,
                            'second_value_type': 'percent',
                            'arrow_type': 'straight'}
        dashboard_items.append(dashboard_item)

        # Наценка

        # dashboard_item = {'header': "Наценка",
        #                     'anchor': markup_anchor,
        #                     'value_main': metrics.markup,
        #                     'prev_main': prev_metrics.markup,
        #                     'main_value_type': 'percent',
        #                     'value_second': None,
        #                     'prev_second': None,
        #                     'second_value_type': 'percent',
        #                     'arrow_type': 'straight'}
        # dashboard_items.append(dashboard_item)

        # Логистика

        dashboard_item = {'header': "Логистика",
                            'anchor': logistic_anchor,
                            'value_main': metrics.logistics,
                            'prev_main': prev_metrics.logistics,
                            'main_value_type': 'currency',
                            'value_second': None,
                            'prev_second': None,
                            'second_value_type': 'percent',
                            'arrow_type': 'back'}
        dashboard_items.append(dashboard_item)

        # Хранение
        dashboard_item = {'header': "Платная приемка / Хранение",
                          'anchor': storage_anchor,
                          'value_main': metrics.acceptance,
                          'prev_main': prev_metrics.acceptance,
                          'main_value_type': 'currency',
                          'value_second': metrics.storage,
                          'prev_second': prev_metrics.storage,
                          'second_value_type': 'currency',
                          'arrow_type': 'back'}
        dashboard_items.append(dashboard_item)

        # # Платная приемка
        # dashboard_item = {'header': "Платная приемка",
        #                   'anchor': acceptance_anchor,
        #                   'value_main': metrics.acceptance,
        #                   'prev_main': prev_metrics.acceptance,
        #                   'main_value_type': 'currency',
        #                   'value_second': None,
        #                   'prev_second': None,
        #                   'second_value_type': 'percent',
        #                   'arrow_type': 'back'}
        # dashboard_items.append(dashboard_item)

        # Прочие удержания и штрафы
        dashboard_item = {'header': "Прочие удержания и штрафы",
                          'anchor': other_deductions_anchor,
                          'value_main': metrics.other_deductions,
                          'prev_main': prev_metrics.other_deductions,
                          'main_value_type': 'currency',
                          'value_second': None,
                          'prev_second': None,
                          'second_value_type': 'percent',
                          'arrow_type': 'back'}
        dashboard_items.append(dashboard_item)


        # Товары без движения

        # dashboard_item = {'header': "Товары без продаж за период",
        #                     'anchor': slow_moving_anchor,
        #                   'value_main': metrics.slow_moving_goods_sum,
        #                   'prev_main': prev_metrics.slow_moving_goods_sum,
        #                   'main_value_type': 'currency',
        #                   'value_second': metrics.slow_moving_goods_quantity,
        #                   'prev_second': prev_metrics.slow_moving_goods_quantity,
        #                   'second_value_type': 'units',
        #                   'arrow_type': 'back'}
        # dashboard_items.append(dashboard_item)

        # Налоги
        dashboard_item = {'header': "Налоги",
                            'anchor': tax_costs_anchor,
                            'value_main': -metrics.tax_costs,
                            'prev_main': -prev_metrics.tax_costs,
                            'main_value_type': 'currency',
                            'value_second': None,
                            'prev_second': None,
                            'second_value_type': 'percent',
                            'arrow_type': 'back'}
        dashboard_items.append(dashboard_item)

        # Остатки по себестоимости

        dashboard_item = {'header': "Остатки по себестоимости",
                          'anchor': goods_at_cost_anchor,
                          'value_main': metrics.stocks_total_at_cost,
                          'prev_main': prev_metrics.stocks_total_at_cost,
                          'main_value_type': 'currency',
                          'value_second': metrics.stocks_quantity,
                          'prev_second': prev_metrics.stocks_quantity,
                          'second_value_type': 'units',
                          'arrow_type': 'straight'}
        dashboard_items.append(dashboard_item)

        # Остатки по цене продажи

        dashboard_item = {'header': "Остатки по цене продажи",
                          'anchor': goods_at_sell_price_anchor,
                          'value_main': metrics.stocks_total_at_sell_price,
                          'prev_main': prev_metrics.stocks_total_at_sell_price,
                          'main_value_type': 'currency',
                          'value_second': None,
                          'prev_second': None,
                          'second_value_type': 'units',
                          'arrow_type': 'straight'}
        dashboard_items.append(dashboard_item)

        # # Оборачиваемость

        dashboard_item = {'header': "Оборачиваемость",
                          'anchor': turnover_anchor,
                          'value_main': metrics.turnover_days,
                          'prev_main': prev_metrics.turnover_days,
                          'main_value_type': 'days',
                          'value_second': None,
                          'prev_second': None,
                          'second_value_type': 'units',
                          'arrow_type': 'back'}
        dashboard_items.append(dashboard_item)

        for item in dashboard_items:
            img = await generate_dashboard_item(header=item['header'],
                                                value_main=item['value_main'],
                                                prev_main=item['prev_main'],
                                                main_value_type=item['main_value_type'],
                                                value_second=item['value_second'],
                                                prev_second=item['prev_second'],
                                                second_value_type=item['second_value_type'],
                                                arrow_type=item['arrow_type'])
            ws.add_image(img, item['anchor'])

        # Форматирование
        thin_color = Side(border_style="thin", color="BFBFBF")
        lightgreenFill = PatternFill(start_color='F4F7ED',
                                     end_color='F4F7ED',
                                     fill_type='solid')
        lightorangeFill = PatternFill(start_color='FDE9D9',
                                      end_color='FDE9D9',
                                      fill_type='solid')
        thin = Side(border_style="thin", color="000000")

        # Форматируем шапку:

        ft = Font(name='Tahoma', size=14, bold=False)
        for row in ws["A1:E6"]:
            for cell in row:
                cell.font = ft

        ft_bold = Font(name='Tahoma', size=14, bold=True)
        for row in ws["A4:E4"]:
            for cell in row:
                cell.font = ft_bold



        # Закрашиваем фон
        lighgreyFill = PatternFill(start_color='F2F2F2',
                                      end_color='F2F2F2',
                                      fill_type='solid')
        thin_grey = Side(border_style="thin", color="F2F2F2")
        for row in ws["A1:T100"]:
            for cell in row:
                cell.fill = lighgreyFill
                cell.border = Border(left=thin_grey,
                                     right=thin_grey,
                                     bottom=thin_grey,
                                     top=thin_grey)

        # Добавляем ячейку, чтобы было красиво с телефона
        invisible_ft = Font(name='Tahoma', size=11, color='F2F2F2', bold=False)
        ws['T57'] = "0"
        empty_cell =ws['T57']
        empty_cell.font = invisible_ft
        empty_cell.fill = lighgreyFill

        # excel_buffer = io.BytesIO()
        # fin_report.save(f"Пример_недельный отчет.xlsx")

        return ws

    except Exception as e:
        # Запись ошибки в лог
        await send_message_to_admin(f'Ошибка в функции создания дэшборда:'
                                    f'\nSeller_id: {seller_id}'
                                    f'\nОшибка: {e}')
        logging.exception("An error occurred: %s", exc_info=e)

